"""Edit sessions: working-copy isolation inside the Office app sandbox containers.

The original file is never written until an explicit, gated save. The working
copy lives inside the target app's sandbox container so that AppleScript-driven
rendering never triggers macOS file-access permission dialogs.
"""
import json
import os
import shutil
import time
import uuid
from pathlib import Path
from typing import Any, Dict, Optional

from docx import Document as WordDocument
from openpyxl import load_workbook

WORD_CONTAINER = Path.home() / "Library/Containers/com.microsoft.Word/Data/tmp"
EXCEL_CONTAINER = Path.home() / "Library/Containers/com.microsoft.Excel/Data/tmp"
FALLBACK_ROOT = Path.home() / ".office_agent/sessions"

_EXT_TO_TYPE = {".docx": "word", ".xlsx": "excel", ".xlsm": "excel"}


def work_root_for(doc_type: str) -> tuple:
    """(root, sandboxed): the app's sandbox container when present (prompt-free
    rendering), else a private fallback where rendering WILL prompt."""
    container = WORD_CONTAINER if doc_type == "word" else EXCEL_CONTAINER
    if container.parent.exists():  # app container present
        return container / "office_agent", True
    return FALLBACK_ROOT, False


class EditSession:
    """One open document: working copy, in-memory object, audit log."""

    def __init__(self, file_path: str):
        original = Path(file_path).expanduser().resolve()
        if not original.exists():
            raise FileNotFoundError(f"Document not found: {original}")
        ext = original.suffix.lower()
        if ext not in _EXT_TO_TYPE:
            raise ValueError(
                f"Unsupported file type '{ext}': expected one of {sorted(_EXT_TO_TYPE)}."
            )
        self.doc_type = _EXT_TO_TYPE[ext]
        self.original_path = original
        self.doc_id = uuid.uuid4().hex[:8]
        root, self.sandboxed = work_root_for(self.doc_type)
        if not self.sandboxed:
            import sys

            print(
                f"[office-agent] 警告：{self.doc_type} 应用容器不存在，工作目录退回 "
                f"{FALLBACK_ROOT} —— 渲染时可能出现 macOS 文件访问弹窗。"
                "运行 `office-agent doctor` 检查。",
                file=sys.stderr,
            )
        self.session_dir = root / self.doc_id
        self.session_dir.mkdir(parents=True, exist_ok=True)
        # Session dirs hold full document contents + audit logs: owner-only.
        os.chmod(self.session_dir, 0o700)
        if not self.sandboxed:
            os.chmod(root, 0o700)
        self.working_path = self.session_dir / f"working{ext}"
        shutil.copy2(original, self.working_path)
        self.preservation_warnings = self._detect_lossy_parts()
        self.doc: Any = self._load()
        self.audit_path = self.session_dir / "audit.jsonl"
        self.dirty = False  # in-memory object differs from working_path on disk
        self.written_paths: set = set()  # files this session has saved to

    def _detect_lossy_parts(self) -> list:
        """openpyxl's reader drops charts/images from pre-existing workbooks.
        Detect them up front so the data loss is disclosed, never silent."""
        if self.doc_type != "excel":
            return []
        import zipfile

        warnings = []
        try:
            names = zipfile.ZipFile(self.working_path).namelist()
        except Exception:
            return []
        n_charts = sum(1 for n in names if n.startswith("xl/charts/chart"))
        n_media = sum(1 for n in names if n.startswith("xl/media/"))
        if n_charts:
            warnings.append(
                f"⚠️ This workbook contains {n_charts} pre-existing chart(s) "
                "that the current editing engine CANNOT preserve — they will "
                "be LOST in any saved output. Inform the user before saving; "
                "do not overwrite the original."
            )
        if n_media:
            warnings.append(
                f"⚠️ This workbook contains {n_media} embedded image(s) that "
                "will be LOST in any saved output. Inform the user before "
                "saving."
            )
        return warnings

    def _load(self) -> Any:
        if self.doc_type == "word":
            return WordDocument(str(self.working_path))
        # keep_vba preserves macros in .xlsm files across load/save.
        return load_workbook(
            str(self.working_path),
            keep_vba=self.working_path.suffix.lower() == ".xlsm",
        )

    def flush(self) -> Path:
        """Write the in-memory object to the working copy on disk."""
        self.doc.save(str(self.working_path))
        self.dirty = False
        return self.working_path

    def reload_from_bytes(self, data: bytes) -> None:
        """Replace in-memory object (and working copy) from snapshot bytes."""
        self.working_path.write_bytes(data)
        self.doc = self._load()
        self.dirty = False

    def to_bytes(self) -> bytes:
        import io

        buf = io.BytesIO()
        self.doc.save(buf)
        return buf.getvalue()

    def save_to(self, dest: Optional[str] = None, overwrite: bool = False) -> Path:
        """Persist to a real destination.

        Overwriting ANY pre-existing file (the original or an unrelated one)
        requires overwrite=True — except files this session already wrote,
        which may be re-saved freely.
        """
        target = Path(dest).expanduser().resolve() if dest else self.original_path
        if target == self.original_path and not overwrite:
            raise PermissionError(
                f"Refusing to overwrite the original file {target} without explicit "
                "approval. Save to a new path, or pass overwrite=True after user consent."
            )
        if (
            target.exists()
            and str(target) not in self.written_paths
            and target != self.original_path
            and not overwrite
        ):
            raise PermissionError(
                f"Refusing to overwrite the existing file {target}: it was not "
                "created by this session. Choose a new path, or pass "
                "overwrite=True after user consent."
            )
        self.flush()
        target.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(self.working_path, target)
        self.written_paths.add(str(target))
        return target

    def default_output_path(self) -> Path:
        p = self.original_path
        return p.with_name(f"{p.stem}.edited{p.suffix}")

    def audit(self, event: str, detail: Dict[str, Any]) -> None:
        record = {"ts": time.time(), "event": event, **detail}
        with self.audit_path.open("a", encoding="utf-8") as f:
            f.write(json.dumps(record, ensure_ascii=False, default=str) + "\n")

    def cleanup(self) -> None:
        shutil.rmtree(self.session_dir, ignore_errors=True)


class SessionManager:
    """Registry of open documents for one agent process."""

    def __init__(self) -> None:
        self.sessions: Dict[str, EditSession] = {}

    def open(self, file_path: str) -> EditSession:
        session = EditSession(file_path)
        self.sessions[session.doc_id] = session
        return session

    def get(self, doc_id: str) -> EditSession:
        if doc_id not in self.sessions:
            raise KeyError(
                f"Unknown doc_id '{doc_id}': open documents are {list(self.sessions)}. "
                "Call open_document first."
            )
        return self.sessions[doc_id]

    def close_all(self, keep_dirs: bool = False) -> None:
        for session in self.sessions.values():
            if not keep_dirs:
                session.cleanup()
        self.sessions.clear()
